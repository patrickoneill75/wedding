"""Builds tests/fixtures/sample_wedding.xlsx: the same layout as the real planning
spreadsheet (Data, Constants, Guest List tabs) with made-up guests, so the importer
can be tested without committing anyone's name or address to this public repo.

Run: python3 scripts/make_test_fixture.py   (needs openpyxl)
"""
from pathlib import Path

from openpyxl import Workbook

OUT = Path(__file__).resolve().parent.parent / "tests" / "fixtures" / "sample_wedding.xlsx"

wb = Workbook()
wb.active.title = "Summary"

data = wb.create_sheet("Data")
data["D1"], data["J1"] = "Estimated", "Actual"
data.append(["Cateogy", "Subcategory", "Item", "Unit Cost", "Units", "Subtotal", "Gratuity %", "Gratuity", "Total Cost", "Actual Cost"])
rows = [
    ("Ceremony & Recption", "Attire", "Dress", None, 1, 0),
    ("Ceremony & Recption", "Entertainment", "DJ (guess)", 1000, 1, 0),
    ("Ceremony & Recption", "Food & Bev", "Yats", "=Constants!C59", 6, 0.2),
    ("Ceremony & Recption", "Food & Bev", "Red Wine", 12, "=Constants!F44", 0.08),
    ("Ceremony & Recption", "Food & Bev", "Non-alcoholic drinks", 0.5, "=Constants!C55", 0),
    ("Communication", "Stationery", "Save the Date Printing", None, "=Constants!C5", 0),
    ("Communication", "Stationery", "Thank You Note Postage", 0.73, "=Constants!C22", 0),
]
for i, (cat, sub, item, unit, units, grat) in enumerate(rows, start=3):
    data.append([cat, sub, item, unit, units, f"=D{i}*E{i}", grat, f"=D{i}*E{i}*G{i}", f"=F{i}+H{i}"])
# A typed-in total with an actual cost, like the real sheet's Save the Date row.
data["I8"] = 294.01
data["J8"] = 294.01

const = wb.create_sheet("Constants")
values = {
    "C18": 0.9, "C19": 0.6, "C20": 0.4,
    "C26": 5, "C27": 15, "C28": 12, "C35": 5, "C37": 2, "C38": 1,
    "C43": 0.4, "C44": 0.5, "C45": 0.5, "C47": 0.6, "C49": 0, "C53": 0.25, "C59": 150,
}
for ref, value in values.items():
    const[ref] = value

guests = wb.create_sheet("Guest List")
guests.append(["Original Order", "Household ID", "First", "Last", "New household", "Ceremony", "Column 7",
               "Bride/Groom", "Priority", "Wedding Party?", "Immediate Family", "Extended", "Full Name(s)",
               "First Name(s)", "Address 1", "Address 2", "City", "State", "Zip", "Country", "Column 2", "Column 3"])
people = [
    (1, "H-001", "Alex", "Example", "x", "x", "Groom", "High", 0, 1, 0, "Alex and Sam Example", "Alex and Sam", "1 Test St", None, "Springfield", "IN", 2134),
    (2, "H-001", "Sam", "Example", None, "x", "Groom", "High", 0, 1, 0, None, None, None, None, None, None, None),
    (3, "H-002", "Jordan ", "Sample", "x", None, "Bride", "Medium", 1, 0, 0, "Jordan Sample", "Jordan", "2 Fake Ave", "Apt 3", "Shelbyville", "IN", 46142),
    # How the real sheet lists a plus-one: its own row, named after the host.
    (4, "H-002", "Guest", "Guest Jordan Sample", None, None, "Bride", "Medium", 0, 0, 0, None, None, None, None, None, None, None),
    (5, "H-003", "Riley", "O'Placeholder", "x", None, "Bride", "High", 0, 0, 1, "Riley O'Placeholder", "Riley", "3 Demo Rd", None, "Ogdenville", "IN", 46203),
]
for p in people:
    guests.append(list(p[:6]) + [None] + list(p[6:]))  # "Column 7" is empty in the real sheet

wb.create_sheet("Round 2 Invites").append(["Original Order", "Household ID", "First", "Last"])
wb.create_sheet("Decor").append(["Categroy", "Item", "Note"])

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"wrote {OUT}")
